from __future__ import annotations

import os
import sys
import queue
import threading
import traceback
from dataclasses import dataclass
from pathlib import Path
from typing import Iterable

import cv2
import numpy as np
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from PIL import Image, ImageTk
import tkinter as tk
from tkinter import filedialog, messagebox, ttk

APP_TITLE = "OMR Sheet Scanner"
CANVAS_W, CANVAS_H = 2480, 3507
IMAGE_EXTENSIONS = {".jpg", ".jpeg", ".png", ".bmp", ".tif", ".tiff"}

# Bubble centers measured on the supplied answer-sheet template (2480 x 3507).
QUESTION_X = {
    0: [414, 506, 596, 686],
    1: [908, 998, 1088, 1178],
    2: [1394, 1484, 1574, 1664],
    3: [1880, 1970, 2060, 2150],
}
QUESTION_Y0 = 2126
QUESTION_DY = 73.2


@dataclass
class ScanResult:
    filename: str
    student_name: str = ""
    school_name: str = ""
    sr_no: str = ""
    grade: str = ""
    district: str = ""
    area: str = ""
    paper: str = ""
    crp_no: str = ""
    school_code: str = ""
    student_gender: str = ""
    answers: list[str] | None = None
    status: str = "OK"


class OMRScanner:
    def __init__(self) -> None:
        self.radius = 16

    @staticmethod
    def _order_points(pts: np.ndarray) -> np.ndarray:
        pts = pts.astype("float32")
        s = pts.sum(axis=1)
        d = np.diff(pts, axis=1).ravel()
        return np.array([
            pts[np.argmin(s)],
            pts[np.argmin(d)],
            pts[np.argmax(s)],
            pts[np.argmax(d)],
        ], dtype="float32")

    def _find_registration_squares(self, gray: np.ndarray) -> list[tuple[float, float, float]]:
        blur = cv2.GaussianBlur(gray, (5, 5), 0)
        bw = cv2.threshold(blur, 90, 255, cv2.THRESH_BINARY_INV)[1]
        contours, _ = cv2.findContours(bw, cv2.RETR_EXTERNAL, cv2.CHAIN_APPROX_SIMPLE)
        h, w = gray.shape
        found: list[tuple[float, float, float]] = []
        for c in contours:
            x, y, cw, ch = cv2.boundingRect(c)
            area = cv2.contourArea(c)
            if not (0.004 * w < cw < 0.035 * w and 0.004 * h < ch < 0.035 * h):
                continue
            ratio = cw / max(ch, 1)
            fill = area / max(cw * ch, 1)
            if 0.72 <= ratio <= 1.28 and fill > 0.55:
                found.append((x + cw / 2, y + ch / 2, cw * ch))
        return found

    def normalize_sheet(self, image: np.ndarray) -> np.ndarray:
        """Correct rotation/perspective and map the sheet to the supplied template size."""
        if image is None:
            raise ValueError("Could not open image")

        # Portrait orientation first.
        if image.shape[1] > image.shape[0]:
            image = cv2.rotate(image, cv2.ROTATE_90_CLOCKWISE)

        gray = cv2.cvtColor(image, cv2.COLOR_BGR2GRAY)
        squares = self._find_registration_squares(gray)
        h, w = gray.shape

        if len(squares) >= 4:
            centers = np.array([[x, y] for x, y, _ in squares], dtype=np.float32)
            # Select the detected square nearest each image corner.
            target_corners = np.array([[0, 0], [w, 0], [w, h], [0, h]], dtype=np.float32)
            chosen = []
            used: set[int] = set()
            for tc in target_corners:
                order = np.argsort(np.linalg.norm(centers - tc, axis=1))
                idx = next(int(i) for i in order if int(i) not in used)
                used.add(idx)
                chosen.append(centers[idx])
            src = self._order_points(np.array(chosen))

            # Registration-square centers in the supplied reference image.
            dst = np.array([
                [174, 198], [2304, 200], [2302, 3278], [166, 3276]
            ], dtype=np.float32)
            matrix = cv2.getPerspectiveTransform(src, dst)
            warped = cv2.warpPerspective(image, matrix, (CANVAS_W, CANVAS_H), borderValue=(255, 255, 255))
        else:
            warped = cv2.resize(image, (CANVAS_W, CANVAS_H), interpolation=cv2.INTER_AREA)

        # Detect upside-down sheets from the two extra side registration squares.
        g2 = cv2.cvtColor(warped, cv2.COLOR_BGR2GRAY)
        sq2 = self._find_registration_squares(g2)
        side_y = [y for x, y, _ in sq2 if x < 280 or x > CANVAS_W - 280]
        middle_y = [y for y in side_y if 650 < y < 2850]
        if middle_y and float(np.median(middle_y)) > CANVAS_H / 2:
            warped = cv2.rotate(warped, cv2.ROTATE_180)

        return warped

    @staticmethod
    def _binary(image: np.ndarray) -> np.ndarray:
        gray = cv2.cvtColor(image, cv2.COLOR_BGR2GRAY)
        gray = cv2.GaussianBlur(gray, (3, 3), 0)
        return cv2.threshold(gray, 0, 255, cv2.THRESH_BINARY_INV + cv2.THRESH_OTSU)[1]

    def _fill_ratio(self, binary: np.ndarray, x: float, y: float, radius: int | None = None) -> float:
        r = radius or self.radius
        x, y = int(round(x)), int(round(y))
        y1, y2 = max(0, y-r), min(binary.shape[0], y+r+1)
        x1, x2 = max(0, x-r), min(binary.shape[1], x+r+1)
        roi = binary[y1:y2, x1:x2]
        yy, xx = np.ogrid[y1-y:y2-y, x1-x:x2-x]
        mask = xx * xx + yy * yy <= r * r
        if roi.size == 0 or mask.sum() == 0:
            return 0.0
        return float(np.mean(roi[mask] > 0))

    def _read_group(self, binary: np.ndarray, centers: Iterable[tuple[float, float]], labels: list[str]) -> str:
        scores = np.array([self._fill_ratio(binary, x, y) for x, y in centers], dtype=float)
        if len(scores) == 0:
            return ""
        order = np.argsort(scores)[::-1]
        best = float(scores[order[0]])
        baseline = float(np.median(scores))
        if best < 0.40 or best - baseline < 0.18:
            return "BLANK"
        selected = [i for i, score in enumerate(scores) if score >= 0.52 and best - score <= 0.12]
        if len(selected) > 1:
            return "MULTIPLE:" + "/".join(labels[i] for i in selected)
        return labels[int(order[0])]

    def scan(self, normalized: np.ndarray, filename: str) -> ScanResult:
        b = self._binary(normalized)

        # Two-digit Sr No.
        digit_x = [418 + 90 * i for i in range(10)]
        sr1 = self._read_group(b, [(x, 778) for x in digit_x], [str(i) for i in range(10)])
        sr2 = self._read_group(b, [(x, 854) for x in digit_x], [str(i) for i in range(10)])
        sr_no = "" if "BLANK" in (sr1, sr2) else sr1 + sr2

        grade = self._read_group(
            b,
            [(1550, 830), (1698, 830), (1838, 830), (1986, 830), (2110, 830), (2190, 830)],
            ["3", "4", "5", "6", "7", "8"],
        )

        district_x = [472 + 90 * i for i in range(10)]
        d1 = self._read_group(b, [(x, 1030) for x in district_x], [str(i) for i in range(10)])
        d2 = self._read_group(b, [(x, 1108) for x in district_x], [str(i) for i in range(10)])
        district = "" if "BLANK" in (d1, d2) else d1 + d2

        area = self._read_group(b, [(1818, 1110), (1986, 1110)], ["Urban", "Rural"])
        paper = self._read_group(
            b,
            [(466, 1338), (616, 1338), (768, 1338), (922, 1338),
             (1074, 1338), (1256, 1338), (1406, 1338), (1550, 1338)],
            ["English", "Urdu", "Math", "Islamiyat", "GK/Sci", "Computer", "Pak.St", "TQM"],
        )
        crp_no = self._read_group(
            b, [(454 + 90 * i, 1544) for i in range(10)],
            [str(i) for i in range(1, 11)],
        )

        school_x = [538 + 90 * i for i in range(10)]
        school_digits = []
        for y in (1718, 1794, 1870):
            school_digits.append(self._read_group(b, [(x, y) for x in school_x], [str(i) for i in range(10)]))
        school_code = "" if any(v == "BLANK" for v in school_digits) else "".join(school_digits)

        gender = self._read_group(b, [(1914, 1878), (2078, 1878)], ["Male", "Female"])

        answers: list[str] = []
        for q in range(1, 61):
            block = (q - 1) // 15
            row = (q - 1) % 15
            y = QUESTION_Y0 + QUESTION_DY * row
            x_values = QUESTION_X[block]
            answers.append(self._read_group(b, [(x, y) for x in x_values], ["A", "B", "C", "D"]))

        return ScanResult(
            filename=filename,
            sr_no=sr_no,
            grade=grade,
            district=district,
            area=area,
            paper=paper,
            crp_no=crp_no,
            school_code=school_code,
            student_gender=gender,
            answers=answers,
        )



def load_answer_key(answer_key_file: Path) -> tuple[list[str], Path]:
    """Load a 60-question A-D answer key from the Excel file selected in the GUI."""
    if not answer_key_file.exists():
        raise FileNotFoundError("The selected answer-key Excel file does not exist.")
    if answer_key_file.suffix.lower() not in {".xlsx", ".xlsm"}:
        raise ValueError("Please select an Excel answer key (.xlsx or .xlsm).")

    wb = load_workbook(answer_key_file, data_only=True, read_only=True)
    try:
        ws = wb.active
        keyed: dict[int, str] = {}
        for row in ws.iter_rows(values_only=True):
            if not row or len(row) < 2:
                continue
            try:
                question_no = int(str(row[0]).strip())
            except (TypeError, ValueError):
                continue
            answer = str(row[1]).strip().upper() if row[1] is not None else ""
            if 1 <= question_no <= 60 and answer in {"A", "B", "C", "D"}:
                keyed[question_no] = answer

        missing = [q for q in range(1, 61) if q not in keyed]
        if missing:
            raise ValueError(
                "The answer key must contain valid A-D answers for Questions 1-60. "
                f"Missing/invalid questions: {missing}"
            )
        return [keyed[q] for q in range(1, 61)], answer_key_file
    finally:
        wb.close()


def calculate_score(answers: list[str] | None, answer_key: list[str]) -> tuple[int, int, float]:
    total_marks = len(answer_key)
    student_answers = answers or []
    obtained_marks = sum(
        1 for index, correct_answer in enumerate(answer_key)
        if index < len(student_answers) and student_answers[index].strip().upper() == correct_answer
    )
    percentage = (obtained_marks / total_marks * 100.0) if total_marks else 0.0
    return total_marks, obtained_marks, percentage


def answer_flags(answers: list[str] | None, answer_key: list[str]) -> list[str]:
    """Return T/F for every question. Only one exact A-D match is True."""
    student_answers = answers or []
    return [
        "T" if index < len(student_answers)
        and student_answers[index].strip().upper() == correct_answer
        else "F"
        for index, correct_answer in enumerate(answer_key)
    ]


def write_excel(results: list[ScanResult], output_file: Path, answer_key: list[str]) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "OMR Results"

    question_headings: list[str] = []
    for question_no in range(1, 61):
        question_headings.extend([f"Q{question_no}", f"Answer {question_no}"])

    headings = [
        "Image File", "Student Name", "School Name", "SR.No.", "Grade", "District",
        "Area", "Paper", "CRP No.", "School Code", "Student Gender",
    ] + question_headings + ["Total Marks", "Obtained Marks", "Percentage", "Status"]
    ws.append(headings)

    for item in results:
        total_marks, obtained_marks, percentage = calculate_score(item.answers, answer_key)
        student_answers = list(item.answers or [])
        student_answers.extend([""] * (60 - len(student_answers)))
        flags = answer_flags(student_answers, answer_key)

        question_values: list[object] = []
        for index in range(60):
            question_values.extend([student_answers[index], flags[index]])

        ws.append([
            item.filename, item.student_name, item.school_name, item.sr_no, item.grade,
            item.district, item.area, item.paper, item.crp_no, item.school_code,
            item.student_gender,
            *question_values, total_marks, obtained_marks, percentage, item.status,
        ])

    header_fill = PatternFill("solid", fgColor="1F4E78")
    true_fill = PatternFill("solid", fgColor="C6EFCE")
    false_fill = PatternFill("solid", fgColor="FFC7CE")
    for cell in ws[1]:
        cell.font = Font(color="FFFFFF", bold=True)
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    ws.row_dimensions[1].height = 34

    widths = {1: 18, 2: 24, 3: 24, 4: 10, 5: 10, 6: 10, 7: 12, 8: 14, 9: 10, 10: 14, 11: 16}
    for col, width in widths.items():
        ws.column_dimensions[get_column_letter(col)].width = width

    first_question_col = 12
    for question_no in range(60):
        q_col = first_question_col + question_no * 2
        answer_col = q_col + 1
        ws.column_dimensions[get_column_letter(q_col)].width = 6
        ws.column_dimensions[get_column_letter(answer_col)].width = 10

        for row_no in range(2, ws.max_row + 1):
            q_cell = ws.cell(row=row_no, column=q_col)
            answer_cell = ws.cell(row=row_no, column=answer_col)
            q_cell.alignment = Alignment(horizontal="center")
            answer_cell.alignment = Alignment(horizontal="center")
            if answer_cell.value == "T":
                answer_cell.fill = true_fill
            elif answer_cell.value == "F":
                answer_cell.fill = false_fill

    total_col = first_question_col + 120
    obtained_col = total_col + 1
    percentage_col = total_col + 2
    status_col = total_col + 3
    ws.column_dimensions[get_column_letter(total_col)].width = 12
    ws.column_dimensions[get_column_letter(obtained_col)].width = 15
    ws.column_dimensions[get_column_letter(percentage_col)].width = 12
    ws.column_dimensions[get_column_letter(status_col)].width = 24

    for row_no in range(2, ws.max_row + 1):
        ws.cell(row=row_no, column=total_col).alignment = Alignment(horizontal="center")
        ws.cell(row=row_no, column=obtained_col).alignment = Alignment(horizontal="center")
        percentage_cell = ws.cell(row=row_no, column=percentage_col)
        percentage_cell.number_format = '0.00"%"'
        percentage_cell.alignment = Alignment(horizontal="center")

    output_file.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output_file)


class ScannerGUI(tk.Tk):
    def __init__(self) -> None:
        super().__init__()
        self.title(APP_TITLE)
        self.geometry("920x650")
        self.minsize(820, 560)
        self.scanner = OMRScanner()
        self.msg_queue: queue.Queue[tuple[str, object]] = queue.Queue()
        self.preview_photo: ImageTk.PhotoImage | None = None

        base = Path(sys.executable).parent if getattr(sys, "frozen", False) else Path(__file__).resolve().parent
        self.default_source = base / "source"
        self.output_folder = base / "output"
        self.selected_source_folder = self.default_source
        self.images_var = tk.StringVar(value=str(self.selected_source_folder))
        default_key = self.default_source / "answer_key.xlsx"
        self.answer_key_var = tk.StringVar(value=str(default_key) if default_key.exists() else "")
        self.output_var = tk.StringVar(value=str(self.output_folder))
        self.status_var = tk.StringVar(value="Ready")

        self._build_ui()
        self.after(150, self._poll_queue)

    def _build_ui(self) -> None:
        pad = {"padx": 10, "pady": 7}
        top = ttk.LabelFrame(self, text="Scanning Inputs")
        top.pack(fill="x", padx=12, pady=10)

        ttk.Label(top, text="Source folder:").grid(row=0, column=0, sticky="w", **pad)
        ttk.Entry(top, textvariable=self.images_var, state="readonly").grid(row=0, column=1, sticky="ew", **pad)
        ttk.Button(top, text="Browse Source Folder", command=self._choose_source_folder).grid(row=0, column=2, **pad)

        ttk.Label(top, text="Answer key:").grid(row=1, column=0, sticky="w", **pad)
        ttk.Entry(top, textvariable=self.answer_key_var).grid(row=1, column=1, sticky="ew", **pad)
        ttk.Button(top, text="Browse Answer Key", command=self._choose_answer_key).grid(row=1, column=2, **pad)

        ttk.Label(top, text="Output folder:").grid(row=2, column=0, sticky="w", **pad)
        ttk.Entry(top, textvariable=self.output_var, state="readonly").grid(row=2, column=1, sticky="ew", **pad)
        ttk.Label(top, text="Automatic").grid(row=2, column=2, **pad)
        top.columnconfigure(1, weight=1)

        action = ttk.Frame(self)
        action.pack(fill="x", padx=12)
        self.start_btn = ttk.Button(action, text="Start Scanning", command=self._start)
        self.start_btn.pack(side="left")
        self.progress = ttk.Progressbar(action, mode="determinate")
        self.progress.pack(side="left", fill="x", expand=True, padx=12)
        ttk.Label(action, textvariable=self.status_var).pack(side="right")

        body = ttk.Panedwindow(self, orient="horizontal")
        body.pack(fill="both", expand=True, padx=12, pady=10)

        left = ttk.LabelFrame(body, text="Processed Preview")
        right = ttk.LabelFrame(body, text="Processing Log")
        body.add(left, weight=1)
        body.add(right, weight=1)

        self.preview = ttk.Label(left, anchor="center")
        self.preview.pack(fill="both", expand=True, padx=6, pady=6)

        self.log = tk.Text(right, wrap="word", state="disabled")
        scroll = ttk.Scrollbar(right, command=self.log.yview)
        self.log.configure(yscrollcommand=scroll.set)
        self.log.pack(side="left", fill="both", expand=True, padx=(6, 0), pady=6)
        scroll.pack(side="right", fill="y", padx=(0, 6), pady=6)

        ttk.Label(
            self,
            text="Multiple, blank, invalid, or incorrect bubbles receive 0 marks. Handwritten names remain blank for manual verification.",
            anchor="w",
        ).pack(fill="x", padx=14, pady=(0, 10))

    def _choose_source_folder(self) -> None:
        path = filedialog.askdirectory(
            title="Select Source Folder Containing OMR Images",
            initialdir=str(self.selected_source_folder),
            mustexist=True,
        )
        if path:
            self.selected_source_folder = Path(path)
            self.images_var.set(str(self.selected_source_folder))
            image_count = sum(
                1 for item in self.selected_source_folder.rglob("*")
                if item.is_file() and item.suffix.lower() in IMAGE_EXTENSIONS
            )
            self._append_log(
                f"Selected source folder: {self.selected_source_folder} "
                f"({image_count} supported image(s))"
            )

    def _choose_answer_key(self) -> None:
        current = Path(self.answer_key_var.get()) if self.answer_key_var.get() else self.default_source
        initial_dir = current.parent if current.suffix else current
        path = filedialog.askopenfilename(
            title="Select Answer Key",
            initialdir=str(initial_dir),
            filetypes=[("Excel workbook", "*.xlsx *.xlsm"), ("All files", "*.*")],
        )
        if path:
            self.answer_key_var.set(path)

    def _append_log(self, text: str) -> None:
        self.log.configure(state="normal")
        self.log.insert("end", text + "\n")
        self.log.see("end")
        self.log.configure(state="disabled")

    def _start(self) -> None:
        source_folder = self.selected_source_folder
        files = sorted(
            p for p in source_folder.rglob("*")
            if p.is_file() and p.suffix.lower() in IMAGE_EXTENSIONS
        ) if source_folder.exists() else []
        answer_key_path = Path(self.answer_key_var.get().strip()) if self.answer_key_var.get().strip() else None
        output = self.output_folder

        if not source_folder.exists() or not source_folder.is_dir():
            messagebox.showwarning(APP_TITLE, "Please click Browse Source Folder and select the folder containing OMR images.")
            return
        if not files:
            messagebox.showwarning(APP_TITLE, "No supported images were found in the selected source folder.")
            return
        if answer_key_path is None:
            messagebox.showwarning(APP_TITLE, "Please click Browse Answer Key and select the Excel answer key.")
            return
        if not answer_key_path.exists():
            messagebox.showerror(APP_TITLE, "The selected answer-key file does not exist.")
            return

        output.mkdir(parents=True, exist_ok=True)
        self.progress["maximum"] = len(files)
        self.progress["value"] = 0
        self.start_btn.configure(state="disabled")
        self.status_var.set(f"0 / {len(files)}")
        self._append_log(f"Starting scan of {len(files)} image(s).")
        threading.Thread(
            target=self._worker,
            args=(files, answer_key_path, output),
            daemon=True,
        ).start()

    def _worker(self, files: list[Path], answer_key_path: Path, output: Path) -> None:
        results: list[ScanResult] = []
        processed_dir = output / "processed_images"
        processed_dir.mkdir(parents=True, exist_ok=True)

        try:
            answer_key, loaded_path = load_answer_key(answer_key_path)
            self.msg_queue.put(("log", f"Answer key loaded: {loaded_path.name} (60 questions)"))
        except Exception as exc:
            self.msg_queue.put(("fatal_message", f"Answer key error:\n{exc}"))
            return

        for index, path in enumerate(files, start=1):
            try:
                image = cv2.imread(str(path))
                normalized = self.scanner.normalize_sheet(image)
                result = self.scanner.scan(normalized, path.name)
                cv2.imwrite(str(processed_dir / path.name), normalized)
                results.append(result)
                _, obtained, percentage = calculate_score(result.answers, answer_key)
                self.msg_queue.put(("log", f"{index}. {path.name}: {obtained}/60 ({percentage:.2f}%)"))
                self.msg_queue.put(("preview", normalized))
            except Exception as exc:
                results.append(ScanResult(filename=path.name, answers=[""] * 60, status=f"ERROR: {exc}"))
                self.msg_queue.put(("log", f"{index}. {path.name}: ERROR - {exc}"))
            self.msg_queue.put(("progress", (index, len(files))))

        excel_path = output / "OMR_Scan_Results.xlsx"
        try:
            write_excel(results, excel_path, answer_key)
            self.msg_queue.put(("done", str(excel_path)))
        except Exception:
            self.msg_queue.put(("fatal", traceback.format_exc()))

    def _poll_queue(self) -> None:
        try:
            while True:
                kind, payload = self.msg_queue.get_nowait()
                if kind == "log":
                    self._append_log(str(payload))
                elif kind == "progress":
                    current, total = payload  # type: ignore[misc]
                    self.progress["value"] = current
                    self.status_var.set(f"{current} / {total}")
                elif kind == "preview":
                    arr = payload  # type: ignore[assignment]
                    rgb = cv2.cvtColor(arr, cv2.COLOR_BGR2RGB)
                    im = Image.fromarray(rgb)
                    im.thumbnail((400, 480))
                    self.preview_photo = ImageTk.PhotoImage(im)
                    self.preview.configure(image=self.preview_photo)
                elif kind == "done":
                    self.start_btn.configure(state="normal")
                    self.status_var.set("Completed")
                    self._append_log(f"Excel created: {payload}")
                    messagebox.showinfo(APP_TITLE, f"Scanning completed.\n\nExcel file:\n{payload}")
                elif kind == "fatal_message":
                    self.start_btn.configure(state="normal")
                    self.status_var.set("Failed")
                    self._append_log(str(payload))
                    messagebox.showerror(APP_TITLE, str(payload))
                elif kind == "fatal":
                    self.start_btn.configure(state="normal")
                    self.status_var.set("Failed")
                    self._append_log(str(payload))
                    messagebox.showerror(APP_TITLE, "Could not create the output Excel file.")
        except queue.Empty:
            pass
        self.after(150, self._poll_queue)


if __name__ == "__main__":
    ScannerGUI().mainloop()
